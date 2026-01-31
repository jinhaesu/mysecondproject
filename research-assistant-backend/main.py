import os
import json
import io
from typing import Optional
from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from dotenv import load_dotenv
from openai import OpenAI
from supabase import create_client, Client
from google.oauth2 import service_account
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload
import PyPDF2
from docx import Document
from pptx import Presentation
from openpyxl import load_workbook
import uvicorn

load_dotenv()

app = FastAPI(title="Research Assistant Backend")

# CORS 설정
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# 환경 변수
SUPABASE_URL = os.getenv("SUPABASE_URL")
SUPABASE_SERVICE_KEY = os.getenv("SUPABASE_SERVICE_KEY")
OPENAI_API_KEY = os.getenv("OPENAI_API_KEY")
GOOGLE_SERVICE_ACCOUNT_KEY = os.getenv("GOOGLE_SERVICE_ACCOUNT_KEY")
GOOGLE_DRIVE_FOLDER_ID = os.getenv("GOOGLE_DRIVE_FOLDER_ID")

# 클라이언트 초기화
openai_client = OpenAI(api_key=OPENAI_API_KEY)
supabase: Client = create_client(SUPABASE_URL, SUPABASE_SERVICE_KEY) if SUPABASE_URL and SUPABASE_SERVICE_KEY else None

def get_drive_service():
    """Google Drive API 서비스 생성"""
    if not GOOGLE_SERVICE_ACCOUNT_KEY:
        raise HTTPException(status_code=500, detail="Google Service Account Key not configured")

    key_data = json.loads(GOOGLE_SERVICE_ACCOUNT_KEY)
    credentials = service_account.Credentials.from_service_account_info(
        key_data,
        scopes=['https://www.googleapis.com/auth/drive.readonly']
    )
    return build('drive', 'v3', credentials=credentials)


def extract_text_from_pdf(content: bytes) -> str:
    """PDF에서 텍스트 추출"""
    reader = PyPDF2.PdfReader(io.BytesIO(content))
    text = ""
    for page in reader.pages:
        text += page.extract_text() or ""
    return text


def extract_text_from_docx(content: bytes) -> str:
    """DOCX에서 텍스트 추출"""
    doc = Document(io.BytesIO(content))
    return "\n".join([para.text for para in doc.paragraphs])


def extract_text_from_pptx(content: bytes) -> str:
    """PPTX에서 텍스트 추출"""
    prs = Presentation(io.BytesIO(content))
    text = ""
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text += shape.text + "\n"
    return text


def extract_text_from_xlsx(content: bytes) -> str:
    """XLSX에서 텍스트 추출"""
    wb = load_workbook(io.BytesIO(content))
    text = ""
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            row_text = " | ".join([str(cell) for cell in row if cell])
            if row_text:
                text += row_text + "\n"
    return text


def get_google_doc_content(service, file_id: str) -> str:
    """Google Docs 내용 추출"""
    request = service.files().export_media(fileId=file_id, mimeType='text/plain')
    content = request.execute()
    return content.decode('utf-8')


def get_google_sheet_content(service, file_id: str) -> str:
    """Google Sheets 내용 추출"""
    request = service.files().export_media(fileId=file_id, mimeType='text/csv')
    content = request.execute()
    return content.decode('utf-8')


def get_google_slides_content(service, file_id: str) -> str:
    """Google Slides 내용 추출"""
    request = service.files().export_media(fileId=file_id, mimeType='text/plain')
    content = request.execute()
    return content.decode('utf-8')


def download_file(service, file_id: str) -> bytes:
    """파일 다운로드 (공유 드라이브 지원)"""
    request = service.files().get_media(fileId=file_id, supportsAllDrives=True)
    fh = io.BytesIO()
    downloader = MediaIoBaseDownload(fh, request)
    done = False
    while not done:
        _, done = downloader.next_chunk()
    fh.seek(0)
    return fh.read()


def create_embedding(text: str) -> list:
    """OpenAI 임베딩 생성"""
    response = openai_client.embeddings.create(
        model="text-embedding-3-small",
        input=text[:8000]  # 토큰 제한
    )
    return response.data[0].embedding


def chunk_text(text: str, chunk_size: int = 1000, overlap: int = 200) -> list:
    """텍스트를 청크로 분할"""
    chunks = []
    start = 0
    while start < len(text):
        end = start + chunk_size
        chunk = text[start:end]
        if chunk.strip():
            chunks.append(chunk)
        start = end - overlap
    return chunks


class ChatRequest(BaseModel):
    query: str
    top_k: int = 5


class ChatResponse(BaseModel):
    answer: str
    sources: list


@app.get("/")
async def health_check():
    return {"status": "healthy", "service": "Research Assistant Backend"}


@app.post("/sync")
async def sync_documents():
    """Google Drive 폴더의 문서들을 동기화"""
    if not supabase:
        raise HTTPException(status_code=500, detail="Supabase not configured")

    try:
        service = get_drive_service()

        # 폴더 내 파일 목록 조회 (공유 드라이브 지원)
        results = service.files().list(
            q=f"'{GOOGLE_DRIVE_FOLDER_ID}' in parents and trashed=false",
            fields="files(id, name, mimeType, modifiedTime)",
            supportsAllDrives=True,
            includeItemsFromAllDrives=True
        ).execute()

        files = results.get('files', [])
        synced = []

        for file in files:
            file_id = file['id']
            file_name = file['name']
            mime_type = file['mimeType']

            try:
                # 파일 유형별 텍스트 추출
                if mime_type == 'application/vnd.google-apps.document':
                    text = get_google_doc_content(service, file_id)
                elif mime_type == 'application/vnd.google-apps.spreadsheet':
                    text = get_google_sheet_content(service, file_id)
                elif mime_type == 'application/vnd.google-apps.presentation':
                    text = get_google_slides_content(service, file_id)
                elif mime_type == 'application/pdf':
                    content = download_file(service, file_id)
                    text = extract_text_from_pdf(content)
                elif mime_type == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document':
                    content = download_file(service, file_id)
                    text = extract_text_from_docx(content)
                elif mime_type == 'application/vnd.openxmlformats-officedocument.presentationml.presentation':
                    content = download_file(service, file_id)
                    text = extract_text_from_pptx(content)
                elif mime_type == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet':
                    content = download_file(service, file_id)
                    text = extract_text_from_xlsx(content)
                else:
                    continue  # 지원하지 않는 파일 형식

                if not text.strip():
                    continue

                # 기존 문서 삭제
                supabase.table('documents').delete().eq('metadata->>file_id', file_id).execute()

                # 텍스트 청킹 및 임베딩
                chunks = chunk_text(text)
                for i, chunk in enumerate(chunks):
                    embedding = create_embedding(chunk)

                    supabase.table('documents').insert({
                        'content': chunk,
                        'metadata': {
                            'file_id': file_id,
                            'file_name': file_name,
                            'mime_type': mime_type,
                            'chunk_index': i,
                            'total_chunks': len(chunks)
                        },
                        'embedding': embedding
                    }).execute()

                synced.append({
                    'name': file_name,
                    'chunks': len(chunks)
                })

            except Exception as e:
                print(f"Error processing {file_name}: {e}")
                continue

        return {
            "success": True,
            "synced_count": len(synced),
            "files": synced
        }

    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/chat", response_model=ChatResponse)
async def chat(request: ChatRequest):
    """RAG 기반 채팅"""
    if not supabase:
        raise HTTPException(status_code=500, detail="Supabase not configured")

    try:
        # 쿼리 임베딩 생성
        query_embedding = create_embedding(request.query)

        # 유사 문서 검색 (Supabase RPC 함수 호출)
        result = supabase.rpc('match_documents', {
            'query_embedding': query_embedding,
            'match_threshold': 0.5,
            'match_count': request.top_k
        }).execute()

        documents = result.data if result.data else []

        # 컨텍스트 구성
        context = "\n\n---\n\n".join([
            f"[출처: {doc['metadata']['file_name']}]\n{doc['content']}"
            for doc in documents
        ])

        # GPT로 답변 생성
        system_prompt = """당신은 생명과학, 화학, 식품공학 분야의 연구 어시스턴트입니다.
제공된 문서 내용을 바탕으로 질문에 답변해주세요.
답변은 한국어로 작성하고, 문서에 없는 내용은 추측하지 마세요.
참고한 문서가 있다면 출처를 명시해주세요."""

        response = openai_client.chat.completions.create(
            model="gpt-4o-mini",
            messages=[
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": f"참고 문서:\n{context}\n\n질문: {request.query}"}
            ],
            temperature=0.3
        )

        answer = response.choices[0].message.content

        # 출처 목록
        sources = list(set([doc['metadata']['file_name'] for doc in documents]))

        return ChatResponse(answer=answer, sources=sources)

    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.get("/documents")
async def list_documents():
    """동기화된 문서 목록 조회"""
    if not supabase:
        raise HTTPException(status_code=500, detail="Supabase not configured")

    try:
        result = supabase.table('documents').select('metadata').execute()

        # 고유 파일 목록 추출
        files = {}
        for doc in result.data:
            file_name = doc['metadata']['file_name']
            if file_name not in files:
                files[file_name] = {
                    'name': file_name,
                    'mime_type': doc['metadata']['mime_type'],
                    'chunks': doc['metadata']['total_chunks']
                }

        return {"documents": list(files.values())}

    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


if __name__ == "__main__":
    port = int(os.getenv("PORT", 8000))
    uvicorn.run(app, host="0.0.0.0", port=port)
